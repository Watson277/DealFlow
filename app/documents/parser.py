from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from uuid import UUID

from docx import Document as DocxDocument

from app.core.config import Settings
from app.core.exceptions import DocumentProcessingError, UnsupportedDocumentError
from app.documents.pdf.mineru import MinerUPDFParser, PDFDocument
from app.documents.pdf.models import DocumentIR


@dataclass(frozen=True, slots=True)
class ParsedDocument:
    text: str
    page_count: int | None
    pdf: PDFDocument | None = None
    document_ir: DocumentIR | None = None


class DocumentParser:
    def __init__(self, settings: Settings | None = None) -> None:
        self.pdf_parser = MinerUPDFParser(settings or Settings())

    @classmethod
    def from_settings(cls, settings: Settings) -> "DocumentParser":
        return cls(settings)

    def parse(
        self,
        content: bytes,
        filename: str,
        document_id: UUID | str | None = None,
    ) -> ParsedDocument:
        extension = Path(filename).suffix.lower()
        try:
            if extension == ".pdf":
                parsed = self._parse_pdf(content, filename, document_id)
            elif extension == ".docx":
                parsed = self._parse_docx(content)
            elif extension in {".md", ".markdown"}:
                parsed = self._parse_markdown(content)
            else:
                raise UnsupportedDocumentError("only PDF, DOCX, and Markdown files are supported")
        except (UnsupportedDocumentError, DocumentProcessingError):
            raise
        except Exception as exc:
            raise DocumentProcessingError(
                f"failed to parse {extension.removeprefix('.').upper()} document"
            ) from exc

        if not parsed.text.strip():
            raise DocumentProcessingError("document contains no extractable text")
        return parsed

    def _parse_pdf(
        self,
        content: bytes,
        filename: str,
        document_id: UUID | str | None,
    ) -> ParsedDocument:
        parsed = self.pdf_parser.parse(content, filename, document_id)
        return ParsedDocument(
            text=parsed.text,
            page_count=parsed.page_count,
            pdf=parsed,
            document_ir=parsed.ir,
        )

    @staticmethod
    def _parse_docx(content: bytes) -> ParsedDocument:
        document = DocxDocument(BytesIO(content))
        blocks: list[str] = []
        blocks.extend(paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text)
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    blocks.append("\t".join(cells))
        return ParsedDocument(text="\n".join(blocks), page_count=None)

    @staticmethod
    def _parse_markdown(content: bytes) -> ParsedDocument:
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise DocumentProcessingError("Markdown files must use UTF-8 encoding") from exc
        if "\x00" in text or any(
            ord(character) < 32 and character not in "\n\r\t" for character in text
        ):
            raise DocumentProcessingError("Markdown file contains unsupported control characters")
        normalized = text.replace("\r\n", "\n").replace("\r", "\n").strip()
        return ParsedDocument(text=normalized, page_count=None)
