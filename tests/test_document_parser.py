from io import BytesIO

import pymupdf
import pytest
from docx import Document as DocxDocument

from app.core.exceptions import DocumentProcessingError, UnsupportedDocumentError
from app.services.document_parser import DocumentParser


def test_parse_pdf_extracts_text_and_page_count() -> None:
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "DealFlow PDF requirement")
    content = document.tobytes()
    document.close()

    parsed = DocumentParser().parse(content, "request.pdf")

    assert "DealFlow PDF requirement" in parsed.text
    assert parsed.page_count == 1


def test_parse_docx_extracts_paragraphs_and_tables() -> None:
    document = DocxDocument()
    document.add_paragraph("DealFlow DOCX requirement")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Delivery"
    table.cell(0, 1).text = "30 days"
    buffer = BytesIO()
    document.save(buffer)

    parsed = DocumentParser().parse(buffer.getvalue(), "request.docx")

    assert "DealFlow DOCX requirement" in parsed.text
    assert "Delivery\t30 days" in parsed.text
    assert parsed.page_count is None


def test_parse_rejects_unsupported_extension() -> None:
    with pytest.raises(UnsupportedDocumentError):
        DocumentParser().parse(b"plain text", "request.txt")


def test_parse_rejects_pdf_without_extractable_text() -> None:
    document = pymupdf.open()
    document.new_page()
    content = document.tobytes()
    document.close()

    with pytest.raises(DocumentProcessingError, match="no extractable text"):
        DocumentParser().parse(content, "empty.pdf")
